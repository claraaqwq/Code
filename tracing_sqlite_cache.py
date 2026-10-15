"""Reusable, bounded-batch Excel imports and indexed account queries."""
import hashlib
import json
import re
import sqlite3
import tempfile
import uuid
from contextlib import closing
from datetime import date, datetime, time, timedelta
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

CACHE_VERSION = 2
CACHE_ROOT = None
BATCH_SIZE = 2000


class OperationCancelled(Exception):
    pass


def check_cancelled(cancel_event):
    if cancel_event is not None and cancel_event.is_set():
        raise OperationCancelled("Operation cancelled.")


def set_cache_directory(directory):
    global CACHE_ROOT
    root = Path(directory).expanduser().resolve()
    if not root.is_dir():
        raise ValueError("Select an existing cache folder.")
    with tempfile.TemporaryFile(dir=root):
        pass
    CACHE_ROOT = root
    return root


def normalize_single_key(value):
    if value is None:
        return ""
    return re.sub(r"^(-?\d+)\.0+$", r"\1", str(value).strip().upper())


def file_signature(path):
    source = Path(path)
    stat = source.stat()
    return hashlib.sha256(
        f"{source.resolve()}|{stat.st_size}|{stat.st_mtime_ns}".encode("utf-8")
    ).hexdigest()


def cache_path_for(path, sheet, header_row, gla_column):
    if CACHE_ROOT is None:
        raise ValueError("Select a cache folder before processing files.")
    key = f"{file_signature(path)}|{sheet}|{header_row}|{gla_column}|{CACHE_VERSION}"
    return CACHE_ROOT / (hashlib.sha256(key.encode("utf-8")).hexdigest() + ".sqlite3")


def _encode(value):
    if isinstance(value, timedelta):
        return {"__excel_type__": "timedelta", "value": value.total_seconds()}
    if isinstance(value, (datetime, date, time)):
        return {"__excel_type__": type(value).__name__, "value": value.isoformat()}
    raise TypeError(f"Unsupported Excel value: {type(value).__name__}")


def _decode(value):
    if value.get("__excel_type__") == "timedelta":
        return timedelta(seconds=value["value"])
    types = {"datetime": datetime, "date": date, "time": time}
    kind = value.get("__excel_type__")
    return types[kind].fromisoformat(value["value"]) if kind in types else value


def _metadata(database):
    # Never create an empty database when checking for an existing import.
    with closing(sqlite3.connect(database.as_uri() + "?mode=ro", uri=True)) as connection:
        result = connection.execute("SELECT value FROM metadata WHERE key='complete'").fetchone()
        if not result:
            raise ValueError("Incomplete SQLite import")
        metadata = json.loads(result[0])
        if metadata["version"] != CACHE_VERSION:
            raise ValueError("Outdated SQLite import")
        return metadata


def load_gla_index(path, sheet, header_row, gla_column):
    """Compatibility name: return small metadata, never a row-number index."""
    database = cache_path_for(path, sheet, header_row, gla_column)
    if not database.exists():
        return None
    try:
        return _metadata(database)
    except (sqlite3.Error, ValueError, KeyError):
        return None


def build_gla_index(path, sheet, header_row, gla_column, progress_callback=None, cancel_event=None, phase_callback=None):
    check_cancelled(cancel_event)
    if phase_callback:
        phase_callback("Opening Excel")
    if header_row < 1:
        raise ValueError("Header rows must be 1 or greater.")
    signature = file_signature(path)
    database = cache_path_for(path, sheet, header_row, gla_column)
    CACHE_ROOT.mkdir(parents=True, exist_ok=True)
    temporary = database.with_name(database.name + "." + uuid.uuid4().hex + ".tmp")
    workbook = None
    connection = None
    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
        worksheet = workbook[sheet]
        rows = worksheet.iter_rows(min_row=header_row, values_only=True)
        raw = next(rows, None)
        if raw is None:
            raise ValueError("Header row is empty.")
        raw = ["" if value is None else str(value).strip() for value in raw]
        if gla_column not in raw:
            raise ValueError(f"Missing column: {gla_column}")
        gla_position = raw.index(gla_column)
        headers, used = [], set()
        for position, name in enumerate(raw, 1):
            base = name or f"Unnamed_{position}"
            candidate, suffix = base, 2
            while candidate in used:
                candidate = f"{base}_{suffix}"
                suffix += 1
            used.add(candidate)
            headers.append(candidate)
        connection = sqlite3.connect(temporary)
        connection.execute("PRAGMA cache_size=-16384")
        connection.execute("PRAGMA temp_store=FILE")
        connection.execute("CREATE TABLE records (row_number INTEGER PRIMARY KEY, gla TEXT NOT NULL, payload TEXT NOT NULL)")
        connection.execute("CREATE TABLE metadata (key TEXT PRIMARY KEY, value TEXT NOT NULL)")
        batch, processed, imported = [], 0, 0
        total = max((worksheet.max_row or header_row) - header_row, 1)
        if phase_callback:
            phase_callback("Importing rows into SQLite")
        for processed, values in enumerate(rows, 1):
            check_cancelled(cancel_event)
            values = tuple(values[:len(headers)])
            values += (None,) * (len(headers) - len(values))
            key = normalize_single_key(values[gla_position])
            if key:
                batch.append((header_row + processed, key, json.dumps(values, default=_encode, ensure_ascii=False)))
                imported += 1
            if len(batch) >= BATCH_SIZE:
                connection.executemany("INSERT INTO records VALUES (?, ?, ?)", batch)
                connection.commit()
                batch.clear()
            if progress_callback and processed % 10000 == 0:
                progress_callback(processed, total, 0)
        if batch:
            connection.executemany("INSERT INTO records VALUES (?, ?, ?)", batch)
            connection.commit()
        if progress_callback:
            progress_callback(processed, max(processed, total), 0)
        check_cancelled(cancel_event)
        if phase_callback:
            phase_callback("Creating GLA index")
        connection.set_progress_handler(
            lambda: int(cancel_event is not None and cancel_event.is_set()), 1000)
        try:
            connection.execute("CREATE INDEX records_gla ON records(gla)")
        except sqlite3.OperationalError:
            check_cancelled(cancel_event)
            raise
        check_cancelled(cancel_event)
        if file_signature(path) != signature:
            raise ValueError("Source file changed during import. Please try again.")
        metadata = {"version": CACHE_VERSION, "headers": headers, "rows": imported}
        connection.execute("INSERT INTO metadata VALUES ('complete', ?)", (json.dumps(metadata),))
        connection.commit()
        connection.close()
        connection = None
        workbook.close()
        workbook = None
        check_cancelled(cancel_event)
        temporary.replace(database)
        return metadata
    except sqlite3.OperationalError:
        check_cancelled(cancel_event)
        raise
    finally:
        if connection is not None:
            connection.close()
        if workbook is not None:
            workbook.close()
        temporary.unlink(missing_ok=True)


def get_or_build_gla_index(path, sheet, header_row, gla_column, progress_callback=None, cancel_event=None, phase_callback=None):
    check_cancelled(cancel_event)
    metadata = load_gla_index(path, sheet, header_row, gla_column)
    if metadata is not None:
        return metadata, True
    return build_gla_index(path, sheet, header_row, gla_column, progress_callback, cancel_event, phase_callback), False


def read_excel_filtered_by_gla(path, sheet, header_row, gla_column,
                               selected_columns, target_gla, progress_callback=None,
                               cache_status_callback=None, cancel_event=None, phase_callback=None):
    metadata, hit = get_or_build_gla_index(path, sheet, header_row, gla_column, progress_callback, cancel_event, phase_callback)
    if cache_status_callback:
        cache_status_callback(hit)
    database = cache_path_for(path, sheet, header_row, gla_column)
    with closing(sqlite3.connect(database.as_uri() + "?mode=ro", uri=True)) as connection:
        check_cancelled(cancel_event)
        cursor = connection.execute(
            "SELECT payload FROM records WHERE gla=? ORDER BY row_number",
            (normalize_single_key(target_gla),))
        def records():
            for row in cursor:
                check_cancelled(cancel_event)
                yield json.loads(row[0], object_hook=_decode)
        return pd.DataFrame.from_records(
            records(),
            columns=metadata["headers"])
