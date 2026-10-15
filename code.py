import re
import threading
import queue
import tkinter as tk
from tkinter import ttk, filedialog
from pathlib import Path

import pandas as pd
import ttkbootstrap as tb
from openpyxl import Workbook, load_workbook
from openpyxl.styles import PatternFill, Font, Border, Side
from openpyxl.utils import get_column_letter

APP_TITLE = "Tracing Validator"
DEFAULT_TOLERANCE = 0.01

BLACK = "111111"
LIGHT_GRAY = "E7E6E6"
CYAN = "DDEBF7"
GREEN = "C6EFCE"
RED = "FFC7CE"
BEG_FILL = "FFF2CC"
VAR_FILL = "DDEBF7"
END_FILL = "E2F0D9"
thin_gray = Side(style="thin", color="BFBFBF")


def normalize_single_key(value):
    if value is None:
        return ""
    text = str(value).strip().upper()
    return re.sub(r"^(-?\d+)\.0+$", r"\1", text)


def norm_key(series):
    return series.map(normalize_single_key).astype("string")


def to_number(series):
    s = series.astype("string").fillna("").str.strip()
    s = s.str.replace(r"^\((.*)\)$", r"-\1", regex=True)
    s = s.str.replace(",", "", regex=False)
    s = s.str.replace("$", "", regex=False)
    return pd.to_numeric(s, errors="coerce").fillna(0.0)


def safe_sheet_name(value):
    value = re.sub(r'[:\\/*?\[\]]+', "_", str(value)).strip()
    return (value or "Tracing")[:31]


def safe_filename(value):
    return re.sub(r'[\\/:*?"<>|]+', "_", str(value)).strip() or "Tracing"


def unique_list(values):
    return list(dict.fromkeys(v for v in values if v and str(v).strip()))


def read_headers_fast(path, sheet, header_row):
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb[sheet]
        row = next(ws.iter_rows(min_row=header_row, max_row=header_row, values_only=True))
        return ["" if v is None else str(v).strip() for v in row]
    finally:
        wb.close()



from tracing_sqlite_cache import (
    load_gla_index, build_gla_index, read_excel_filtered_by_gla,
    set_cache_directory,
)


class MappingBox(ttk.LabelFrame):
    def __init__(self, parent, title, roles):
        super().__init__(parent, text=title, padding=10)
        self.vars = {}
        self.combos = {}
        self.labels = {}
        self.role_rows = {}
        for i, role in enumerate(roles):
            label = ttk.Label(self, text=role)
            label.grid(row=i, column=0, sticky="w", padx=(0, 8), pady=4)
            var = tk.StringVar()
            combo = ttk.Combobox(self, textvariable=var, state="normal", width=32)
            combo.grid(row=i, column=1, sticky="ew", pady=4)
            self.vars[role] = var
            self.combos[role] = combo
            self.labels[role] = label
            self.role_rows[role] = i
        self.columnconfigure(1, weight=1)

    def set_columns(self, columns):
        vals = [""] + [str(c) for c in columns if str(c).strip()]
        for combo in self.combos.values():
            combo["values"] = vals

    def get(self, role):
        return self.vars[role].get().strip()

    def show_role(self, role, visible=True):
        if role not in self.vars:
            return
        if visible:
            row = self.role_rows[role]
            self.labels[role].grid(row=row, column=0, sticky="w", padx=(0, 8), pady=4)
            self.combos[role].grid(row=row, column=1, sticky="ew", pady=4)
        else:
            self.labels[role].grid_remove()
            self.combos[role].grid_remove()
            self.vars[role].set("")


class App(tb.Window):
    def __init__(self):
        super().__init__(themename="darkly")
        self.title(APP_TITLE)
        self.geometry("1260x900")
        self.minsize(1100, 500)

        self.style.configure(
            "Accent.TButton", background="#FFE600", foreground="#111111",
            font=("Segoe UI Semibold", 10), padding=(14, 8)
        )
        self.style.map(
            "Accent.TButton",
            background=[("active", "#FFD400")],
            foreground=[("active", "#111111")]
        )
        self.style.configure("Secondary.TButton", font=("Segoe UI", 10), padding=(11, 7))

        self.je_files = []
        self.tb_beg_path = ""
        self.tb_end_path = ""
        self.rf_path = ""
        self.je_sheets = {}
        self.tb_beg_sheet = None
        self.tb_end_sheet = None
        self.rf_sheet = None
        self.je_columns = []
        self.tb_beg_columns = []
        self.tb_end_columns = []
        self.rf_columns = []
        self.ui_queue = queue.Queue()
        self.worker_running = False
        self._build_ui()
        self.protocol("WM_DELETE_WINDOW", self.confirm_close)
        self.after(100, self._process_ui_queue)

    def _build_ui(self):
        viewport = ttk.Frame(self)
        viewport.pack(fill="both", expand=True)
        self.content_canvas = tk.Canvas(viewport, highlightthickness=0)
        scrollbar = ttk.Scrollbar(viewport, orient="vertical", command=self.content_canvas.yview)
        scrollbar.pack(side="right", fill="y")
        self.content_canvas.pack(side="left", fill="both", expand=True)
        self.content_canvas.configure(yscrollcommand=scrollbar.set)
        outer = ttk.Frame(self.content_canvas, padding=18)
        content_window = self.content_canvas.create_window((0, 0), window=outer, anchor="nw")
        outer.bind("<Configure>", lambda event: self.content_canvas.configure(
            scrollregion=self.content_canvas.bbox("all")))
        self.content_canvas.bind("<Configure>", lambda event:
            self.content_canvas.itemconfigure(content_window, width=event.width))
        self.bind("<MouseWheel>", self._scroll_content, add="+")

        header = ttk.Frame(outer)
        header.pack(fill="x", pady=(0, 16))
        bar = tk.Frame(header, bg="#FFE600", width=7, height=58)
        bar.pack(side="left", padx=(0, 14))
        bar.pack_propagate(False)
        titles = ttk.Frame(header)
        titles.pack(side="left")
        ttk.Label(titles, text="TRACING", font=("Segoe UI", 23, "bold")).pack(anchor="w")
        ttk.Label(titles, text="VALIDATOR", foreground="#FFE600",
                  font=("Segoe UI", 12, "bold")).pack(anchor="w")
        ttk.Label(header, text="GLA + BUSINESS UNIT · WORKING PAPER GENERATOR",
                  foreground="#999999", font=("Segoe UI", 9, "bold")).pack(side="right")

        level_box = ttk.LabelFrame(outer, text="01  ROLLFORWARD LEVEL", padding=12)
        level_box.pack(fill="x", pady=(0, 12))
        ttk.Label(
            level_box,
            text="Select the level at which the Rollforward must reconcile:"
        ).pack(anchor="w", pady=(0, 6))
        self.rf_level_var = tk.StringVar(value="GL Account + Business Unit")
        self.rf_level_combo = ttk.Combobox(
            level_box,
            textvariable=self.rf_level_var,
            state="readonly",
            values=[
                "GL Account",
                "GL Account + Business Unit",
                "GL Account + Business Unit + Source",
            ],
            width=44,
        )
        self.rf_level_combo.pack(anchor="w")
        self.rf_level_combo.bind("<<ComboboxSelected>>", self.on_rollforward_level_change)

        files = ttk.LabelFrame(outer, text="02  SOURCE FILES", padding=12)
        files.pack(fill="x", pady=(0, 12))
        files.columnconfigure(1, weight=1)

        self.je_header_var = tk.StringVar(value="31")
        self.tb_beg_header_var = tk.StringVar(value="1")
        self.tb_end_header_var = tk.StringVar(value="1")
        self.rf_header_var = tk.StringVar(value="1")

        def file_row(row, label, cmd, attr, header_var):
            ttk.Label(files, text=label).grid(row=row, column=0, sticky="w", padx=(0, 12), pady=4)
            status = ttk.Label(files, text="Not selected", foreground="#BBBBBB")
            status.grid(row=row, column=1, sticky="w", pady=4)
            setattr(self, attr, status)
            ttk.Label(files, text="Header Row").grid(row=row, column=2, padx=(12, 6), pady=4)
            ttk.Entry(files, textvariable=header_var, width=7).grid(
                row=row, column=3, padx=(0, 12), pady=4)
            ttk.Button(files, text="Browse", style="Secondary.TButton",
                       command=cmd).grid(row=row, column=4, sticky="e", pady=4)

        file_row(0, "Journal Entries (one or many)", self.add_je, "je_label", self.je_header_var)
        file_row(1, "TB Beginning", lambda: self.pick_single("tb_beg"), "tb_beg_label", self.tb_beg_header_var)
        file_row(2, "TB Ending", lambda: self.pick_single("tb_end"), "tb_end_label", self.tb_end_header_var)
        file_row(3, "Rollforward", lambda: self.pick_single("rf"), "rf_label", self.rf_header_var)

        ttk.Label(files, text="Local cache folder").grid(row=4, column=0, sticky="w", pady=(12, 4))
        self.cache_directory_var = tk.StringVar()
        ttk.Entry(files, textvariable=self.cache_directory_var, state="readonly").grid(
            row=4, column=1, columnspan=3, sticky="ew", padx=(0, 12), pady=(12, 4))
        ttk.Button(files, text="Choose folder", style="Secondary.TButton",
                   command=self.choose_cache_directory).grid(row=4, column=4, pady=(12, 4))

        buttons_row = ttk.Frame(files)
        buttons_row.grid(row=5, column=0, columnspan=5, sticky="ew", pady=(10, 0))
        buttons_row.columnconfigure(0, weight=3)
        buttons_row.columnconfigure(1, weight=1)
        ttk.Button(buttons_row, text="LOAD HEADERS", style="Accent.TButton",
                   command=self.load_files).grid(row=0, column=0, sticky="ew", padx=(0, 6))
        ttk.Button(buttons_row, text="BUILD / REFRESH CACHE", style="Secondary.TButton",
                   command=self.build_cache).grid(row=0, column=1, sticky="ew", padx=(6, 0))

        maps = ttk.Frame(outer)
        maps.pack(fill="both", expand=True)

        self.je_map = MappingBox(
            maps, "JOURNAL ENTRIES", ["GLA", "Business Unit", "Source", "Amount"])
        self.tb_beg_map = MappingBox(
            maps, "TB BEGINNING", ["GLA", "Business Unit", "Source", "Amount"])
        self.tb_end_map = MappingBox(
            maps, "TB ENDING", ["GLA", "Business Unit", "Source", "Amount"])
        self.rf_map = MappingBox(
            maps, "ROLLFORWARD",
            ["GLA", "Business Unit", "Source", "Beginning", "Variance", "Ending"])

        self.je_map.grid(row=0, column=0, sticky="nsew", padx=(0, 6), pady=(0, 6))
        self.tb_beg_map.grid(row=0, column=1, sticky="nsew", padx=(6, 0), pady=(0, 6))
        self.tb_end_map.grid(row=1, column=0, sticky="nsew", padx=(0, 6), pady=(6, 0))
        self.rf_map.grid(row=1, column=1, sticky="nsew", padx=(6, 0), pady=(6, 0))
        maps.columnconfigure(0, weight=1)
        maps.columnconfigure(1, weight=1)

        # Manual dimension values are useful when a TB does not physically
        # contain Business Unit and/or Source.
        self.tb_beg_manual_bu_var = tk.StringVar()
        self.tb_beg_manual_source_var = tk.StringVar()
        self.tb_end_manual_bu_var = tk.StringVar()
        self.tb_end_manual_source_var = tk.StringVar()

        self.tb_bu_hints = {}
        for box, manual_var, columns_attr in [
            (self.tb_beg_map, self.tb_beg_manual_bu_var, "tb_beg_columns"),
            (self.tb_end_map, self.tb_end_manual_bu_var, "tb_end_columns"),
        ]:
            for role in ["Source", "Amount"]:
                box.role_rows[role] += 1
                box.labels[role].grid_configure(row=box.role_rows[role])
                box.combos[role].grid_configure(row=box.role_rows[role])
            hint = ttk.Label(box, foreground="#BBBBBB")
            hint.grid(row=2, column=1, sticky="w", pady=(0, 4))
            self.tb_bu_hints[box] = hint
            box.vars["Business Unit"].trace_add(
                "write", lambda *args, b=box, v=manual_var, c=columns_attr:
                self._update_tb_business_unit(b, v, c))
            self._update_tb_business_unit(box, manual_var, columns_attr)
        self.tb_beg_manual_source_label = ttk.Label(self.tb_beg_map, text="Manual Source")
        self.tb_beg_manual_source_entry = ttk.Entry(
            self.tb_beg_map, textvariable=self.tb_beg_manual_source_var, width=32)

        self.tb_end_manual_source_label = ttk.Label(self.tb_end_map, text="Manual Source")
        self.tb_end_manual_source_entry = ttk.Entry(
            self.tb_end_map, textvariable=self.tb_end_manual_source_var, width=32)

        self.tb_beg_manual_source_label.grid(row=5, column=0, sticky="w", pady=4)
        self.tb_beg_manual_source_entry.grid(row=5, column=1, sticky="ew", pady=4)
        self.tb_end_manual_source_label.grid(row=5, column=0, sticky="w", pady=4)
        self.tb_end_manual_source_entry.grid(row=5, column=1, sticky="ew", pady=4)

        self.on_rollforward_level_change()

        run = ttk.LabelFrame(outer, text="04  RUN TRACING", padding=12)
        run.pack(fill="x", pady=(12, 0))
        for c in range(4):
            run.columnconfigure(c, weight=1)

        ttk.Label(run, text="GLA ACCOUNT").grid(row=0, column=0, sticky="w")
        self.account_var = tk.StringVar()
        self.account_combo = ttk.Combobox(run, textvariable=self.account_var, state="normal", width=28)
        self.account_combo.grid(row=1, column=0, padx=(0, 12), pady=(4, 0), sticky="ew")

        ttk.Label(run, text="TOLERANCE").grid(row=0, column=1, sticky="w")
        self.tolerance_var = tk.StringVar(value=str(DEFAULT_TOLERANCE))
        ttk.Entry(run, textvariable=self.tolerance_var, width=14).grid(
            row=1, column=1, padx=(0, 12), pady=(4, 0), sticky="ew")

        ttk.Button(run, text="REFRESH GLA", style="Secondary.TButton",
                   command=self.refresh_accounts).grid(row=1, column=2, padx=(0, 12), sticky="ew")
        self.generate_button = ttk.Button(
            run, text="GENERATE WORKING PAPER", style="Accent.TButton",
            command=self.generate)
        self.generate_button.grid(row=1, column=3, sticky="ew")

        self.progress_var = tk.DoubleVar(value=0)
        self.progress = ttk.Progressbar(run, variable=self.progress_var, maximum=100, mode="determinate")
        self.progress.grid(row=2, column=0, columnspan=4, sticky="ew", pady=(14, 4))
        self.progress_text = ttk.Label(run, text="Ready", foreground="#BBBBBB")
        self.progress_text.grid(row=3, column=0, columnspan=4, sticky="w")

    def active_dimensions(self):
        level = self.rf_level_var.get()
        if level == "GL Account":
            return []
        if level == "GL Account + Business Unit":
            return ["Business Unit"]
        return ["Business Unit", "Source"]

    def active_group_helpers(self):
        return ["__GLA"] + [
            "__BU" if d == "Business Unit" else "__SOURCE"
            for d in self.active_dimensions()
        ]

    def on_rollforward_level_change(self, event=None):
        if not hasattr(self, "je_map"):
            return
        dims = self.active_dimensions()
        use_bu = "Business Unit" in dims
        use_source = "Source" in dims

        for box in [self.je_map, self.tb_beg_map, self.tb_end_map, self.rf_map]:
            box.show_role("Business Unit", use_bu)
            box.show_role("Source", use_source)

        for hint in self.tb_bu_hints.values():
            if use_bu:
                hint.grid()
            else:
                hint.grid_remove()

        if hasattr(self, "tb_beg_manual_source_label"):
            widgets = [
                (self.tb_beg_manual_source_label, self.tb_beg_manual_source_entry, use_source, 5),
                (self.tb_end_manual_source_label, self.tb_end_manual_source_entry, use_source, 5),
            ]
            for label, entry, visible, row in widgets:
                if visible:
                    label.grid(row=row, column=0, sticky="w", pady=4)
                    entry.grid(row=row, column=1, sticky="ew", pady=4)
                else:
                    label.grid_remove()
                    entry.grid_remove()

    def _scroll_content(self, event):
        if event.widget.winfo_toplevel() != self:
            return
        if event.widget.winfo_class() in {"TCombobox", "Text", "Listbox"}:
            return
        bounds = self.content_canvas.bbox("all")
        if bounds and bounds[3] > self.content_canvas.winfo_height() and event.delta:
            steps = -1 if event.delta > 0 else 1
            self.content_canvas.yview_scroll(steps * max(1, abs(event.delta) // 120), "units")
            return "break"

    def _update_tb_business_unit(self, box, manual_var, columns_attr):
        value = box.get("Business Unit")
        is_column = value in getattr(self, columns_attr)
        manual_var.set(value if value and not is_column else "")
        hint = ("Column selected" if is_column else
                "Fixed value for all rows" if value else
                "Select a column or enter a fixed value")
        self.tb_bu_hints[box].configure(text=hint)

    def _apply_progress(self, percentage, text):
        percentage = max(0, min(float(percentage), 100))
        self.progress_var.set(percentage)
        self.progress_text.config(text=text)

    def set_progress(self, percentage, text):
        if threading.current_thread() is threading.main_thread():
            self._apply_progress(percentage, text)
        else:
            self.ui_queue.put(("progress", percentage, text))

    def _process_ui_queue(self):
        try:
            while True:
                item = self.ui_queue.get_nowait()
                kind = item[0]
                if kind == "progress":
                    self._apply_progress(item[1], item[2])
                elif kind == "done":
                    self.worker_running = False
                    self.generate_button.config(state="normal", text="GENERATE WORKING PAPER")
                    self._apply_progress(100, f"Completed: {Path(item[1]).name}")
                    self.show_message("Working paper ready", item[2] + "\n\nSaved to:\n" + item[1])
                elif kind == "error":
                    self.worker_running = False
                    self.generate_button.config(state="normal", text="GENERATE WORKING PAPER")
                    self._apply_progress(0, "Generation failed.")
                    self.show_error("Error", item[1])
                elif kind == "cache_done":
                    self.worker_running = False
                    self.generate_button.config(state="normal", text="GENERATE WORKING PAPER")
                    self._apply_progress(100, "Persistent SQLite cache ready.")
                    self.show_message("Cache ready", item[1])
                elif kind == "cache_error":
                    self.worker_running = False
                    self.generate_button.config(state="normal", text="GENERATE WORKING PAPER")
                    self._apply_progress(0, "Cache build failed.")
                    self.show_error("Cache error", item[1])
        except queue.Empty:
            pass
        finally:
            self.after(100, self._process_ui_queue)


    def add_je(self):
        paths = filedialog.askopenfilenames(
            title="Select one or more JE files",
            filetypes=[("Excel files", "*.xlsx *.xlsm")]
        )
        if paths:
            self.je_files = list(paths)
            names = ", ".join(Path(p).name for p in self.je_files[:3])
            if len(self.je_files) > 3:
                names += "..."
            self.je_label.config(text=f"{len(self.je_files)} file(s): {names}")

    def pick_single(self, kind):
        path = filedialog.askopenfilename(filetypes=[("Excel files", "*.xlsx *.xlsm")])
        if not path:
            return
        setattr(self, f"{kind}_path", path)
        getattr(self, f"{kind}_label").config(text=Path(path).name)

    def _create_dialog(self, title, subtitle, accent="#FFE600"):
        win = tk.Toplevel(self)
        win.withdraw()
        win.title(title)
        win.transient(self)
        win.resizable(False, False)
        tk.Frame(win, bg=accent, height=5).pack(fill="x")
        body = ttk.Frame(win, padding=24)
        body.pack(fill="both", expand=True)
        ttk.Label(body, text="TRACING VALIDATOR", foreground=accent,
                  font=("Segoe UI", 9, "bold")).pack(anchor="w")
        ttk.Label(body, text=title, font=("Segoe UI", 17, "bold"),
                  wraplength=490).pack(anchor="w", pady=(10, 6))
        if subtitle:
            ttk.Label(body, text=subtitle, foreground="#BBBBBB",
                      wraplength=490).pack(anchor="w", pady=(0, 16))
        win.bind("<Escape>", lambda event: win.destroy())
        return win, body

    def _show_dialog(self, win, focus_widget):
        self.update_idletasks()
        win.update_idletasks()
        width, height = win.winfo_reqwidth(), win.winfo_reqheight()
        x = self.winfo_rootx() + (self.winfo_width() - width) // 2
        y = self.winfo_rooty() + (self.winfo_height() - height) // 2
        # Secondary monitors can have negative desktop coordinates. Do not
        # clamp to the primary screen or use Tk's negative-edge geometry syntax.
        if self.tk.call("tk", "windowingsystem") == "win32":
            import ctypes
            from ctypes import wintypes

            get_parent = ctypes.windll.user32.GetParent
            get_parent.argtypes = [wintypes.HWND]
            get_parent.restype = wintypes.HWND
            set_window_pos = ctypes.windll.user32.SetWindowPos
            set_window_pos.argtypes = [wintypes.HWND, wintypes.HWND,
                                      ctypes.c_int, ctypes.c_int, ctypes.c_int,
                                      ctypes.c_int, wintypes.UINT]
            set_window_pos.restype = wintypes.BOOL
            win.geometry(f"{width}x{height}")
            win.update_idletasks()
            hwnd = get_parent(win.winfo_id()) or win.winfo_id()
            # Keep the requested size and stacking order; move in desktop coordinates.
            set_window_pos(hwnd, None, x, y, 0, 0, 0x0001 | 0x0004 | 0x0010)
        else:
            win.geometry(f"{width}x{height}{x:+d}{y:+d}")
        win.deiconify()
        win.wait_visibility()
        win.lift(self)
        win.grab_set()
        focus_widget.focus_set()
        self.wait_window(win)

    def confirm_close(self):
        message = (
            "An operation is still running. Closing the application will interrupt it. "
            "Do you want to close anyway?"
            if self.worker_running else
            "Do you want to close Tracing Validator? Your completed cache files will remain available."
        )
        win, body = self._create_dialog("Close Tracing Validator?", message)
        result = {"close": False}

        def accept():
            result["close"] = True
            win.destroy()

        actions = ttk.Frame(body)
        actions.pack(fill="x", pady=(8, 0))
        cancel = ttk.Button(actions, text="KEEP OPEN", style="Secondary.TButton",
                            command=win.destroy)
        cancel.pack(side="left")
        ttk.Button(actions, text="CLOSE APPLICATION", style="Accent.TButton",
                   command=accept).pack(side="right", padx=(20, 0))
        self._show_dialog(win, cancel)
        if result["close"]:
            self.destroy()

    def show_error(self, title, message):
        self.show_message(title, message, error=True)

    def show_message(self, title, message, error=False):
        accent = "#FF7474" if error else "#FFE600"
        subtitle = "Please review the details below." if error else "Operation completed successfully."
        win, body = self._create_dialog(title, subtitle, accent)
        details = ttk.Frame(body)
        details.pack(fill="both", expand=True, pady=(0, 20))
        message = str(message)
        text = tk.Text(details, width=62, height=min(12, max(3, len(message) // 60 + message.count("\n") + 1)),
                       wrap="word", bg="#292929", fg="#EEEEEE",
                       relief="flat", font=("Segoe UI", 10), padx=12, pady=12)
        scroll = ttk.Scrollbar(details, orient="vertical", command=text.yview)
        scroll.pack(side="right", fill="y")
        text.configure(yscrollcommand=scroll.set)
        text.pack(side="left", fill="both", expand=True)
        text.insert("1.0", message)
        text.configure(state="disabled")
        button = ttk.Button(body, text="CLOSE" if error else "CONTINUE",
                            style="Accent.TButton", command=win.destroy)
        button.pack(anchor="e")
        win.bind("<Return>", lambda event: win.destroy())
        self._show_dialog(win, button)

    def choose_sheet(self, path, title):
        wb = load_workbook(path, read_only=True, data_only=True)
        sheets = wb.sheetnames
        wb.close()
        if len(sheets) == 1:
            return sheets[0]
        win, body = self._create_dialog(title, Path(path).name)
        var = tk.StringVar(value=sheets[0])
        ttk.Label(body, text="Select the worksheet to use").pack(anchor="w", pady=(0, 8))
        combo = ttk.Combobox(body, values=sheets, textvariable=var, state="readonly", width=55)
        combo.pack(fill="x", pady=(0, 22))
        result = {"sheet": None}
        def ok():
            result["sheet"] = var.get()
            win.destroy()
        actions = ttk.Frame(body)
        actions.pack(fill="x")
        ttk.Button(actions, text="CANCEL", style="Secondary.TButton",
                   command=win.destroy).pack(side="left")
        ttk.Button(actions, text="USE WORKSHEET", style="Accent.TButton",
                   command=ok).pack(side="right")
        win.bind("<Return>", lambda event: ok())
        self._show_dialog(win, combo)
        return result["sheet"]

    def load_files(self):
        if not self.je_files or not self.tb_beg_path or not self.tb_end_path or not self.rf_path:
            self.show_error("Missing files",
                                 "Select JE(s), TB Beginning, TB Ending and Rollforward.")
            return
        try:
            je_header = int(self.je_header_var.get().strip())
            beg_header = int(self.tb_beg_header_var.get().strip())
            end_header = int(self.tb_end_header_var.get().strip())
            rf_header = int(self.rf_header_var.get().strip())
            if min(je_header, beg_header, end_header, rf_header) < 1:
                raise ValueError("Header rows must be 1 or greater.")

            expected = None
            total = len(self.je_files) + 3
            done = 0
            self.je_sheets = {}

            for path in self.je_files:
                sheet = self.choose_sheet(path, "JE sheet")
                if not sheet:
                    return
                self.je_sheets[path] = sheet
                cols = read_headers_fast(path, sheet, je_header)
                if expected is None:
                    expected = cols
                done += 1
                self.set_progress(2 + 28 * done / total, f"JE header: {Path(path).name}")

            self.je_columns = expected or []

            self.tb_beg_sheet = self.choose_sheet(self.tb_beg_path, "TB Beginning sheet")
            if not self.tb_beg_sheet: return
            self.tb_beg_columns = read_headers_fast(self.tb_beg_path, self.tb_beg_sheet, beg_header)

            self.tb_end_sheet = self.choose_sheet(self.tb_end_path, "TB Ending sheet")
            if not self.tb_end_sheet: return
            self.tb_end_columns = read_headers_fast(self.tb_end_path, self.tb_end_sheet, end_header)

            self.rf_sheet = self.choose_sheet(self.rf_path, "Rollforward sheet")
            if not self.rf_sheet: return
            self.rf_columns = read_headers_fast(self.rf_path, self.rf_sheet, rf_header)

            self.je_map.set_columns(self.je_columns)
            self.tb_beg_map.set_columns(self.tb_beg_columns)
            self.tb_end_map.set_columns(self.tb_end_columns)
            self.rf_map.set_columns(self.rf_columns)
            self._update_tb_business_unit(
                self.tb_beg_map, self.tb_beg_manual_bu_var, "tb_beg_columns")
            self._update_tb_business_unit(
                self.tb_end_map, self.tb_end_manual_bu_var, "tb_end_columns")
            self.set_progress(30, "Headers loaded. Complete the mapping and enter a GLA.")
            self.show_message("Loaded", "Headers loaded successfully. No data rows were loaded yet.")
        except Exception as exc:
            self.show_error("Load error", str(exc))
            self.set_progress(0, "Header loading failed.")

    def choose_cache_directory(self):
        if self.worker_running:
            self.show_error("Processing in progress", "Wait for the current operation before changing the cache folder.")
            return False
        directory = filedialog.askdirectory(
            parent=self, title="Select a folder for the local data cache",
            initialdir=self.cache_directory_var.get() or None, mustexist=True)
        if not directory:
            return False
        try:
            root = set_cache_directory(directory)
            self.cache_directory_var.set(str(root))
            return True
        except (OSError, ValueError) as exc:
            self.show_error("Cache folder unavailable", str(exc))
            return False

    def ensure_cache_directory(self):
        if not self.cache_directory_var.get():
            return self.choose_cache_directory()
        try:
            set_cache_directory(self.cache_directory_var.get())
            return True
        except (OSError, ValueError) as exc:
            self.show_error("Cache folder unavailable", str(exc))
            return False

    def build_cache(self):
        if self.worker_running:
            return
        try:
            self.validate_mapping()
            if not self.ensure_cache_directory():
                return
            self.worker_config = {
                "je": {r: self.je_map.get(r) for r in ["GLA", "Business Unit", "Amount"]},
                "beg": {r: self.tb_beg_map.get(r) for r in ["GLA", "Business Unit", "Amount"]},
                "end": {r: self.tb_end_map.get(r) for r in ["GLA", "Business Unit", "Amount"]},
                "rf": {r: self.rf_map.get(r) for r in ["GLA", "Business Unit", "Beginning", "Variance", "Ending"]},
                "je_header": int(self.je_header_var.get().strip()),
                "beg_header": int(self.tb_beg_header_var.get().strip()),
                "end_header": int(self.tb_end_header_var.get().strip()),
                "rf_header": int(self.rf_header_var.get().strip()),
            }
            self.worker_running = True
            self.generate_button.config(state="disabled")
            self._apply_progress(1, "Building persistent SQLite cache...")
            threading.Thread(target=self._build_cache_worker, daemon=True).start()
        except Exception as exc:
            self.show_error("Cache", str(exc))

    def _build_cache_worker(self):
        try:
            cfg = self.worker_config
            sources = []
            for i, path in enumerate(self.je_files, 1):
                sources.append((
                    f"JE {i}/{len(self.je_files)}", path, self.je_sheets[path],
                    cfg["je_header"], cfg["je"]["GLA"]
                ))
            sources.extend([
                ("TB Beginning", self.tb_beg_path, self.tb_beg_sheet,
                 cfg["beg_header"], cfg["beg"]["GLA"]),
                ("TB Ending", self.tb_end_path, self.tb_end_sheet,
                 cfg["end_header"], cfg["end"]["GLA"]),
                ("Rollforward", self.rf_path, self.rf_sheet,
                 cfg["rf_header"], cfg["rf"]["GLA"]),
            ])

            total_sources = len(sources)
            cache_hits = 0
            for source_no, (label, path, sheet, header, gla_col) in enumerate(sources):
                def cb(processed, total, unique_accounts, label=label, source_no=source_no):
                    frac = min(processed / max(total, 1), 1)
                    pct = 2 + 93 * ((source_no + frac) / total_sources)
                    self.set_progress(
                        pct,
                        f"{label}: importing {processed:,}/{total:,} rows into SQLite"
                    )

                existing = load_gla_index(path, sheet, header, gla_col)
                if existing is not None:
                    cache_hits += 1
                    self.set_progress(
                        2 + 93 * ((source_no + 1) / total_sources),
                        f"{label}: cache already valid."
                    )
                else:
                    build_gla_index(path, sheet, header, gla_col, cb)

            msg = (
                f"Persistent SQLite cache is ready.\n\n"
                f"Sources: {total_sources}\n"
                f"Already cached: {cache_hits}\n"
                f"Built now: {total_sources - cache_hits}\n\n"
                f"Future tracings with these unchanged files will reuse the cache."
            )
            self.ui_queue.put(("cache_done", msg))
        except Exception as exc:
            self.ui_queue.put(("cache_error", str(exc)))

    def validate_column(self, box, role, columns, source, required=True):
        value = box.get(role)
        if not value:
            if required:
                raise ValueError(f"Select or enter '{role}' for {source}.")
            return
        if value not in columns:
            raise ValueError(f"Column '{value}' was not found in {source}.")

    def validate_mapping(self):
        if not self.je_columns:
            raise ValueError("Load workbook headers first.")

        dims = self.active_dimensions()

        self.validate_column(self.je_map, "GLA", self.je_columns, "JE")
        self.validate_column(self.je_map, "Amount", self.je_columns, "JE")
        for dim in dims:
            self.validate_column(self.je_map, dim, self.je_columns, "JE")

        for box, cols, name, manual_bu, manual_source in [
            (
                self.tb_beg_map, self.tb_beg_columns, "TB Beginning",
                self.tb_beg_manual_bu_var.get().strip(),
                self.tb_beg_manual_source_var.get().strip(),
            ),
            (
                self.tb_end_map, self.tb_end_columns, "TB Ending",
                self.tb_end_manual_bu_var.get().strip(),
                self.tb_end_manual_source_var.get().strip(),
            ),
        ]:
            self.validate_column(box, "GLA", cols, name)
            self.validate_column(box, "Amount", cols, name)
            if "Business Unit" in dims and not manual_bu:
                self.validate_column(box, "Business Unit", cols, name)
            if "Source" in dims and not manual_source:
                self.validate_column(box, "Source", cols, name)

        self.validate_column(self.rf_map, "GLA", self.rf_columns, "Rollforward")
        for dim in dims:
            self.validate_column(self.rf_map, dim, self.rf_columns, "Rollforward")
        for role in ["Beginning", "Variance", "Ending"]:
            self.validate_column(self.rf_map, role, self.rf_columns, "Rollforward")

    def refresh_accounts(self):
        try:
            if not self.rf_sheet:
                raise ValueError("Load workbook headers first.")
            gla_col = self.rf_map.get("GLA")
            if not gla_col or gla_col not in self.rf_columns:
                raise ValueError("Map a valid Rollforward GLA column first.")
            header_row = int(self.rf_header_var.get().strip())
            self.set_progress(5, "Scanning Rollforward GLAs...")
            wb = load_workbook(self.rf_path, read_only=True, data_only=True)
            try:
                ws = wb[self.rf_sheet]
                headers = ["" if v is None else str(v).strip() for v in next(
                    ws.iter_rows(min_row=header_row, max_row=header_row, values_only=True))]
                pos = headers.index(gla_col)
                values = set()
                for row in ws.iter_rows(min_row=header_row + 1, values_only=True):
                    if pos < len(row):
                        v = normalize_single_key(row[pos])
                        if v:
                            values.add(v)
            finally:
                wb.close()
            values = sorted(values)
            self.account_combo["values"] = values
            if values:
                self.account_var.set(values[0])
            self.set_progress(30, f"{len(values):,} GLA account(s) available.")
        except Exception as exc:
            self.show_error("Error", str(exc))

    def _prepared(self, df, mapping, manual_bu=None):
        out = df.copy()
        out["__GLA"] = norm_key(out[mapping.get("GLA")])
        if manual_bu and manual_bu.strip():
            out["__BU"] = normalize_single_key(manual_bu)
        else:
            out["__BU"] = norm_key(out[mapping.get("Business Unit")])
        out["__AMOUNT"] = to_number(out[mapping.get("Amount")])
        return out

    def load_data_for_generation(self, account):
        cfg = self.worker_config
        je_header = cfg["je_header"]
        beg_header = cfg["beg_header"]
        end_header = cfg["end_header"]
        rf_header = cfg["rf_header"]
        total_sources = len(self.je_files) + 3
        source_index = 0

        def callback_for(label, index):
            def callback(processed, total, matched):
                safe_total = total if isinstance(total, (int, float)) and total > 0 else processed
                source_fraction = min(processed / max(safe_total, 1), 1)
                overall = 5 + 65 * ((index + source_fraction) / total_sources)
                self.set_progress(
                    overall,
                    f"{label}: importing {processed:,}/{safe_total:,} rows into SQLite"
                )
            return callback

        def cache_status_for(label, index):
            def callback(cache_hit):
                if cache_hit:
                    overall = 5 + 65 * ((index + 0.25) / total_sources)
                    self.set_progress(overall, f"{label}: persistent cache found; locating account...")
            return callback

        je_parts = []
        for n, path in enumerate(self.je_files, 1):
            part = read_excel_filtered_by_gla(
                path, self.je_sheets[path], je_header, cfg["je"]["GLA"],
                [], account,
                callback_for(f"JE {n}/{len(self.je_files)}", source_index),
                cache_status_for(f"JE {n}/{len(self.je_files)}", source_index)
            )
            part["Source_File"] = Path(path).name
            je_parts.append(part)
            source_index += 1

        je_df = (
            pd.concat(je_parts, ignore_index=True)
            if je_parts else pd.DataFrame()
        )

        beg_df = read_excel_filtered_by_gla(
            self.tb_beg_path, self.tb_beg_sheet, beg_header,
            cfg["beg"]["GLA"], [], account,
            callback_for("TB Beginning", source_index),
            cache_status_for("TB Beginning", source_index)
        )
        source_index += 1

        end_df = read_excel_filtered_by_gla(
            self.tb_end_path, self.tb_end_sheet, end_header,
            cfg["end"]["GLA"], [], account,
            callback_for("TB Ending", source_index),
            cache_status_for("TB Ending", source_index)
        )
        source_index += 1

        rf_df = read_excel_filtered_by_gla(
            self.rf_path, self.rf_sheet, rf_header,
            cfg["rf"]["GLA"], [], account,
            callback_for("Rollforward", source_index),
            cache_status_for("Rollforward", source_index)
        )

        return je_df, beg_df, end_df, rf_df

    def generate(self):
        if self.worker_running:
            return

        try:
            self.set_progress(1, "Validating mapping...")
            self.validate_mapping()
            if not self.ensure_cache_directory():
                return

            account = normalize_single_key(self.account_var.get())
            if not account:
                raise ValueError("Select or enter a GLA account.")

            tolerance = float(self.tolerance_var.get().strip())

            save_path = filedialog.asksaveasfilename(
                defaultextension=".xlsx",
                initialfile=f"{safe_filename(account)}_Tracing.xlsx",
                filetypes=[("Excel workbook", "*.xlsx")]
            )
            if not save_path:
                return

            # Snapshot all Tkinter values before entering the worker thread.
            self.worker_config = {
                "level": self.rf_level_var.get(),
                "dimensions": self.active_dimensions(),
                "je": {r: self.je_map.get(r) for r in ["GLA", "Business Unit", "Source", "Amount"]},
                "beg": {r: self.tb_beg_map.get(r) for r in ["GLA", "Business Unit", "Source", "Amount"]},
                "end": {r: self.tb_end_map.get(r) for r in ["GLA", "Business Unit", "Source", "Amount"]},
                "rf": {r: self.rf_map.get(r) for r in ["GLA", "Business Unit", "Source", "Beginning", "Variance", "Ending"]},
                "beg_manual_bu": self.tb_beg_manual_bu_var.get().strip(),
                "beg_manual_source": self.tb_beg_manual_source_var.get().strip(),
                "end_manual_bu": self.tb_end_manual_bu_var.get().strip(),
                "end_manual_source": self.tb_end_manual_source_var.get().strip(),
                "je_header": int(self.je_header_var.get().strip()),
                "beg_header": int(self.tb_beg_header_var.get().strip()),
                "end_header": int(self.tb_end_header_var.get().strip()),
                "rf_header": int(self.rf_header_var.get().strip()),
            }

            self.worker_running = True
            self.generate_button.config(state="disabled", text="PROCESSING...")
            self._apply_progress(2, "Starting background processing...")

            worker = threading.Thread(
                target=self._generate_worker,
                args=(account, tolerance, save_path),
                daemon=True
            )
            worker.start()

        except Exception as exc:
            self.show_error("Error", str(exc))
            self._apply_progress(0, "Generation failed.")

    def _generate_worker(self, account, tolerance, save_path):
        try:
            je_df, beg_df, end_df, rf_df = self.load_data_for_generation(account)
            self.set_progress(72, "Preparing filtered records...")

            cfg = self.worker_config
            dims = cfg["dimensions"]

            def prepare_detail(df, mapping, manual_bu="", manual_source=""):
                out = df.copy()
                out["__GLA"] = norm_key(out[mapping["GLA"]])

                if "Business Unit" in dims:
                    if manual_bu:
                        out["__BU"] = normalize_single_key(manual_bu)
                    else:
                        out["__BU"] = norm_key(out[mapping["Business Unit"]])

                if "Source" in dims:
                    if manual_source:
                        out["__SOURCE"] = normalize_single_key(manual_source)
                    else:
                        out["__SOURCE"] = norm_key(out[mapping["Source"]])

                out["__AMOUNT"] = to_number(out[mapping["Amount"]])
                return out

            je = prepare_detail(je_df, cfg["je"])
            beg = prepare_detail(
                beg_df, cfg["beg"],
                cfg["beg_manual_bu"], cfg["beg_manual_source"])
            end = prepare_detail(
                end_df, cfg["end"],
                cfg["end_manual_bu"], cfg["end_manual_source"])

            rf = rf_df.copy()
            rf["__GLA"] = norm_key(rf[cfg["rf"]["GLA"]])
            if "Business Unit" in dims:
                rf["__BU"] = norm_key(rf[cfg["rf"]["Business Unit"]])
            if "Source" in dims:
                rf["__SOURCE"] = norm_key(rf[cfg["rf"]["Source"]])
            rf["__BEG"] = to_number(rf[cfg["rf"]["Beginning"]])
            rf["__VAR"] = to_number(rf[cfg["rf"]["Variance"]])
            rf["__END"] = to_number(rf[cfg["rf"]["Ending"]])

            missing = []
            if je.empty: missing.append("JE")
            if beg.empty: missing.append("TB Beginning")
            if end.empty: missing.append("TB Ending")
            if rf.empty: missing.append("Rollforward")
            if missing:
                raise ValueError(
                    f"No records were found for GLA '{account}' in: "
                    + ", ".join(missing)
                    + ". Check the mapped GLA columns and header rows."
                )

            self.set_progress(78, "Calculating validation totals...")

            group_cols = ["__GLA"]
            if "Business Unit" in dims:
                group_cols.append("__BU")
            if "Source" in dims:
                group_cols.append("__SOURCE")

            rf_s = rf.groupby(group_cols, dropna=False).agg(
                RF_Beginning=("__BEG", "sum"),
                RF_Variance=("__VAR", "sum"),
                RF_Ending=("__END", "sum")
            ).reset_index()

            beg_s = beg.groupby(
                group_cols, dropna=False
            )["__AMOUNT"].sum().reset_index(name="TB_Beginning_Total")

            je_s = je.groupby(
                group_cols, dropna=False
            )["__AMOUNT"].sum().reset_index(name="JE_Total")

            end_s = end.groupby(
                group_cols, dropna=False
            )["__AMOUNT"].sum().reset_index(name="TB_Ending_Total")

            val = rf_s.merge(beg_s, on=group_cols, how="outer")
            val = val.merge(je_s, on=group_cols, how="outer")
            val = val.merge(end_s, on=group_cols, how="outer")

            numeric = [
                "RF_Beginning", "RF_Variance", "RF_Ending",
                "TB_Beginning_Total", "JE_Total", "TB_Ending_Total"
            ]
            val[numeric] = val[numeric].fillna(0.0)

            val["Beginning_Difference"] = (
                val["RF_Beginning"] - val["TB_Beginning_Total"])
            val["Beginning_Result"] = (
                val["Beginning_Difference"].abs().le(tolerance)
                .map({True: "PASS", False: "FAIL"}))

            val["Activity_Difference"] = val["RF_Variance"] - val["JE_Total"]
            val["Activity_Result"] = (
                val["Activity_Difference"].abs().le(tolerance)
                .map({True: "PASS", False: "FAIL"}))

            val["Ending_Difference"] = val["RF_Ending"] - val["TB_Ending_Total"]
            val["Ending_Result"] = (
                val["Ending_Difference"].abs().le(tolerance)
                .map({True: "PASS", False: "FAIL"}))

            rename_map = {"__GLA": "GLA"}
            if "Business Unit" in dims:
                rename_map["__BU"] = "Business Unit"
            if "Source" in dims:
                rename_map["__SOURCE"] = "Source"
            val.rename(columns=rename_map, inplace=True)

            helper_cols = ["__GLA", "__BU", "__SOURCE", "__AMOUNT"]
            je_original = je.drop(columns=helper_cols, errors="ignore")
            beg_original = beg.drop(columns=helper_cols, errors="ignore")
            end_original = end.drop(columns=helper_cols, errors="ignore")
            rf_original = rf.drop(
                columns=["__GLA", "__BU", "__SOURCE", "__BEG", "__VAR", "__END"],
                errors="ignore")

            self.set_progress(86, "Writing the Excel working paper...")
            self.write_working_paper(
                save_path, account, rf_original, beg_original,
                end_original, je_original, val, tolerance
            )

            summary = (
                f"Working paper created.\n\n"
                f"Rollforward level: {cfg['level']}\n"
                f"JE rows: {len(je):,}\n"
                f"TB Beginning rows: {len(beg):,}\n"
                f"TB Ending rows: {len(end):,}\n"
                f"Rollforward rows: {len(rf):,}"
            )
            self.ui_queue.put(("done", save_path, summary))

        except Exception as exc:
            self.ui_queue.put(("error", str(exc)))

    def write_working_paper(self, path, account, rf_df, beg_df, end_df, je_df, val_df, tolerance):
        wb = Workbook()
        ws = wb.active
        ws.title = safe_sheet_name(f"Tracing {account}")
        ws.sheet_view.showGridLines = False

        ws["A1"] = "TRACING WORKING PAPER"
        ws["A1"].font = Font(size=18, bold=True)
        ws["A2"] = f"GLA: {account}"
        ws["A3"] = f"Rollforward Level: {self.worker_config['level']}"
        ws["A4"] = f"Tolerance: {tolerance}"

        def section_title(row, text):
            ws.cell(row, 1, text)
            ws.cell(row, 1).font = Font(bold=True, size=12)
            ws.cell(row, 1).fill = PatternFill("solid", fgColor=BLACK)
            ws.cell(row, 1).font = Font(bold=True, size=12, color="FFFFFF")
            return row + 1

        def write_df(start_row, df, header_fill=LIGHT_GRAY):
            headers = list(df.columns)
            for j, h in enumerate(headers, 1):
                cell = ws.cell(start_row, j, h)
                cell.font = Font(bold=True)
                cell.fill = PatternFill("solid", fgColor=header_fill)
                cell.border = Border(bottom=thin_gray)
            for i, values in enumerate(df.itertuples(index=False, name=None), start_row + 1):
                for j, value in enumerate(values, 1):
                    ws.cell(i, j, value)
            return start_row + 1 + len(df), headers

        def add_total(row, header_row, headers, amount_name, df_len):
            if amount_name not in headers:
                return row
            c = headers.index(amount_name) + 1
            ws.cell(row, max(1, c - 1), "TOTAL").font = Font(bold=True)
            value = f"=SUM({get_column_letter(c)}{header_row + 1}:{get_column_letter(c)}{row - 1})" if df_len else 0
            ws.cell(row, c, value)
            ws.cell(row, c).font = Font(bold=True)
            ws.cell(row, c).fill = PatternFill("solid", fgColor=CYAN)
            return row + 1

        row = 6
        row = section_title(row, "ROLLFORWARD — DETAIL")
        rf_header = row
        row, rf_headers = write_df(row, rf_df, "D9EAD3")
        cfg = self.worker_config
        for header, color in [(cfg["rf"]["Beginning"], BEG_FILL),
                              (cfg["rf"]["Variance"], VAR_FILL),
                              (cfg["rf"]["Ending"], END_FILL)]:
            if header in rf_headers:
                c = rf_headers.index(header) + 1
                for r in range(rf_header, row):
                    ws.cell(r, c).fill = PatternFill("solid", fgColor=color)
        row += 2

        row = section_title(row, "VALIDATION BY GLA + BUSINESS UNIT")
        validation_cols = ["GLA"]
        if "Business Unit" in cfg["dimensions"]:
            validation_cols.append("Business Unit")
        if "Source" in cfg["dimensions"]:
            validation_cols.append("Source")
        validation_cols += [
            "RF_Beginning", "TB_Beginning_Total",
            "Beginning_Difference", "Beginning_Result",
            "RF_Variance", "JE_Total",
            "Activity_Difference", "Activity_Result",
            "RF_Ending", "TB_Ending_Total",
            "Ending_Difference", "Ending_Result"
        ]
        vh = row
        row, headers = write_df(row, val_df[validation_cols].copy(), "D9EAF7")
        for result_col in ["Beginning_Result", "Activity_Result", "Ending_Result"]:
            c = headers.index(result_col) + 1
            for r in range(vh + 1, row):
                color = GREEN if ws.cell(r, c).value == "PASS" else RED
                ws.cell(r, c).fill = PatternFill("solid", fgColor=color)
        row += 2

        for title, df, amount in [
            ("TB BEGINNING — DETAIL", beg_df, cfg["beg"]["Amount"]),
            ("TB ENDING — DETAIL", end_df, cfg["end"]["Amount"]),
            ("JOURNAL ENTRIES — DETAIL", je_df, cfg["je"]["Amount"]),
        ]:
            row = section_title(row, title)
            h = row
            row, headers = write_df(row, df, "D9EAF7")
            row = add_total(row, h, headers, amount, len(df))
            row += 2

        for col in range(1, max(ws.max_column, 1) + 1):
            letter = get_column_letter(col)
            max_len = 0
            for r in range(1, min(ws.max_row, 250) + 1):
                value = ws.cell(r, col).value
                if value is not None:
                    max_len = max(max_len, len(str(value)))
            ws.column_dimensions[letter].width = min(max(max_len + 2, 11), 32)

        for cells in ws.iter_rows():
            for cell in cells:
                if isinstance(cell.value, (int, float)) or (
                    isinstance(cell.value, str) and cell.value.startswith("=")
                ):
                    cell.number_format = '#,##0.00;[Red](#,##0.00);-'

        wb.save(path)


if __name__ == "__main__":
    App().mainloop()
